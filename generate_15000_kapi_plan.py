# -*- coding: utf-8 -*-
"""
15.000 Kapı Kapasite Planlaması - Hat Dağılım Planı
Reçeteler:
  - 1 Kapı = 2 Plaka Levha (2.07 m)
  - 1 Kapı = 3.5 Boy Seren (2.10 m)
  - 1 Kapı = 5.0 Boy Kasa  (2.15 m)  ← Kullanıcı reçetesi
  - 1 Kapı = 5.0 Boy Pervaz (2.20 m)
"""

import openpyxl
from openpyxl.styles import (
    Font, Alignment, PatternFill, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

def generate():
    KAPI_TARGET = 15_000

    # ─────────────────────────────────────────────────────────────
    # 1. REÇETE ve TOPLAM İHTİYAÇ
    # ─────────────────────────────────────────────────────────────
    bom = {
        'Levha (Plaka)':  {'factor': 2.0,  'unit': 'Plaka'},
        'Seren (Boy)':    {'factor': 3.5,  'unit': 'Boy'},
        'Kasa (Boy)':     {'factor': 5.0,  'unit': 'Boy'},
        'Pervaz (Boy)':   {'factor': 5.0,  'unit': 'Boy'},
    }

    # ─────────────────────────────────────────────────────────────
    # 2. HATLAR – FİİLİ GÜNLÜK KAPASİTELER (24 saat, ortalama)
    # ─────────────────────────────────────────────────────────────
    # Hat adı, ürün grubu, ürün adı, birim, saatlik adet, günlük (24s), boy uzunluğu
    lines = [
        # HAT 201 – Levha
        ('Hat 201', 'Levha', 'WPC Kapı Levhası (Genel Ort.)', 'Plaka', 41.7, 1000, '2.07 m'),

        # HAT 104 – Seren
        ('Hat 104', 'Seren', 'Masif WPC Kapı Sereni 38×45 mm', 'Boy', 28.2,  676, '2.10 m'),
        # HAT 105 – Seren
        ('Hat 105', 'Seren', 'Masif WPC Kapı Sereni 38×45 mm', 'Boy', 26.7,  640, '2.10 m'),
        # HAT 107 – Seren(39mm) isteğe bağlı
        ('Hat 107', 'Seren', 'Seren 39 mm (Yedek)', 'Boy', 27.9, 670, '2.10 m'),

        # HAT 106 – 140 mm Kasa
        ('Hat 106', 'Kasa', '140 mm Kasa', 'Boy', 14.0, 336, '2.15 m'),
        # HAT 108 – 140 mm Kasa / 100 mm Kasa
        ('Hat 108', 'Kasa', '140 mm Kasa', 'Boy', 14.7, 352, '2.15 m'),
        # HAT 102 – 100 mm Kasa
        ('Hat 102', 'Kasa', '100 mm Kasa', 'Boy', 15.7, 376, '2.15 m'),

        # HAT 109 – 50×80 Pervaz
        ('Hat 109', 'Pervaz', '50×80 Pervaz', 'Boy', 50.2, 1204, '2.20 m'),
        # HAT 103 – 50×80 Pervaz
        ('Hat 103', 'Pervaz', '50×80 Pervaz', 'Boy', 31.5,  756, '2.20 m'),
    ]

    # ─────────────────────────────────────────────────────────────
    # 3. PLAN – HAT GÖREV ATAMALARI (15.000 kapı için)
    # ─────────────────────────────────────────────────────────────
    # Toplam ihtiyaçlar
    total_levha  = KAPI_TARGET * 2.0    # 30 000 plaka
    total_seren  = KAPI_TARGET * 3.5    # 52 500 boy
    total_kasa   = KAPI_TARGET * 5.0    # 75 000 boy
    total_pervaz = KAPI_TARGET * 5.0    # 75 000 boy

    # Strateji:
    #  Levha  → Hat 201 tek hat
    #  Seren  → Hat 104 + Hat 105 primer, Hat 107 üçüncü hat
    #  Kasa   → Hat 106 (140mm) + Hat 108 (140mm) + Hat 102 (100mm/fazla kapasite)
    #  Pervaz → Hat 109 + Hat 103

    # Gün sayısı = toplam / günlük
    plan_tasks = [
        # (Hat, Ürün, Birim, Günlük, Toplam Hedef, Kapsama Oranı)
        # Levha ─────────────────────────────────────────────────
        ('Hat 201', 'WPC Kapı Levhası', 'Plaka', 1000, total_levha,  'Tek Hat – Tüm Levha'),

        # Seren ─────────────────────────────────────────────────
        ('Hat 104', 'Seren 38×45 mm',   'Boy',    676, 0,  'Ana Hat'),
        ('Hat 105', 'Seren 38×45 mm',   'Boy',    640, 0,  'İkinci Hat'),
        ('Hat 107', 'Seren 39 mm',       'Boy',    670, 0,  'Üçüncü Hat (Gerektiğinde)'),

        # Kasa ──────────────────────────────────────────────────
        ('Hat 106', '140 mm Kasa',      'Boy',    336, 0, '140 mm Ana Hat'),
        ('Hat 108', '140 mm Kasa',      'Boy',    352, 0, '140 mm İkinci Hat'),
        ('Hat 102', '100 mm Kasa',      'Boy',    376, 0, '100 mm Hat'),

        # Pervaz ─────────────────────────────────────────────────
        ('Hat 109', '50×80 Pervaz',     'Boy',   1204, 0, 'Ana Hat'),
        ('Hat 103', '50×80 Pervaz',     'Boy',    756, 0, 'İkinci Hat'),
    ]

    # Seren gün payı: Hat 104 + 105 kombine ile önce 2 hat kapatır
    seren_2hat_daily = 676 + 640            # 1316 boy/gün
    seren_3hat_daily = 676 + 640 + 670      # 1986 boy/gün
    seren_days_2hat  = total_seren / seren_2hat_daily   # ~39.9
    seren_days_3hat  = total_seren / seren_3hat_daily   # ~26.4

    kasa_3hat_daily  = 336 + 352 + 376     # 1064 boy/gün
    kasa_days_3hat   = total_kasa / kasa_3hat_daily     # ~70.5

    pervaz_2hat_daily = 1204 + 756          # 1960 boy/gün
    pervaz_days_2hat  = total_pervaz / pervaz_2hat_daily   # ~38.3

    levha_days = total_levha / 1000         # 30.0

    # ─────────────────────────────────────────────────────────────
    # 4. EXCEL OLUŞTURMA
    # ─────────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()

    # ──── RENK PALETİ ─────────────────────────────────────────────
    C_HEADER_DARK  = '1F3864'   # koyu lacivert
    C_HEADER_MED   = '2E75B6'   # orta mavi
    C_LEVHA        = 'FCE4D6'   # açık turuncu
    C_SEREN        = 'E2EFDA'   # açık yeşil
    C_KASA         = 'DDEEFF'   # açık mavi
    C_PERVAZ       = 'FFF2CC'   # açık sarı
    C_TOTAL        = 'F2F2F2'   # gri
    C_ACCENT       = 'C6EFCE'   # yeşil vurgu

    group_colors = {
        'Levha':  C_LEVHA,
        'Seren':  C_SEREN,
        'Kasa':   C_KASA,
        'Pervaz': C_PERVAZ,
    }

    def fill(hex_color):
        return PatternFill('solid', fgColor=hex_color)

    def font(bold=False, size=10, color='000000', italic=False):
        return Font(bold=bold, size=size, color=color, italic=italic)

    def align(h='center', v='center', wrap=False):
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

    def border_thin():
        s = Side(style='thin', color='AAAAAA')
        return Border(left=s, right=s, top=s, bottom=s)

    def border_medium():
        s = Side(style='medium', color='2E75B6')
        return Border(left=s, right=s, top=s, bottom=s)

    # ═══════════════════════════════════════════════════════════════
    # SAYFA 1 – BOM & HAT ÖZET PLANI
    # ═══════════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = '15K Kapasite Planı'
    ws1.sheet_view.showGridLines = False

    # ── Başlık ──
    ws1.merge_cells('A1:L1')
    ws1['A1'] = f'ERGÜNBAŞ WPC – 15.000 KAPI KAPASİTE PLANLAMA MATRİSİ'
    ws1['A1'].font = Font(bold=True, size=16, color='FFFFFF')
    ws1['A1'].fill = fill(C_HEADER_DARK)
    ws1['A1'].alignment = align()

    ws1.merge_cells('A2:L2')
    ws1['A2'] = ('Reçete: 1 Kapı = 2 Plaka Levha | 3.5 Boy Seren | 5 Boy Kasa | 5 Boy Pervaz  '
                 '|  Net 720 dk (12s) Vardiya  |  Sürekli 24 Saat Üretim')
    ws1['A2'].font = Font(italic=True, size=9, color='FFFFFF')
    ws1['A2'].fill = fill(C_HEADER_MED)
    ws1['A2'].alignment = align()

    ws1.row_dimensions[1].height = 30
    ws1.row_dimensions[2].height = 18

    # ── BÖLÜM 1: BOM Tablosu ──────────────────────────────────────
    ws1.merge_cells('A3:L3')
    ws1['A3'] = 'BÖLÜM 1 – REÇETE (BOM) ve TOPLAM İHTİYAÇ'
    ws1['A3'].font = Font(bold=True, size=11, color='FFFFFF')
    ws1['A3'].fill = fill(C_HEADER_MED)
    ws1['A3'].alignment = align('left')
    ws1['A3'].border = border_thin()
    ws1.row_dimensions[3].height = 18

    bom_headers = ['Ürün Grubu', 'Birim', 'Kapı Başı Reçete', '15.000 Kapı için Toplam']
    bom_widths  = [6, 3, 6, 9]
    for ci, h in enumerate(bom_headers, 1):
        c = ws1.cell(4, ci, h)
        c.font = font(bold=True, color='FFFFFF')
        c.fill = fill(C_HEADER_DARK)
        c.alignment = align()
        c.border = border_thin()

    bom_rows = [
        ('Levha (Hat 201)',  'Plaka', '2 Plaka / Kapı',  f'{int(total_levha):,}  Plaka'),
        ('Seren (Hat 104/105/107)', 'Boy',  '3.5 Boy / Kapı', f'{int(total_seren):,}  Boy'),
        ('Kasa  (Hat 106/108/102)',  'Boy',  '5.0 Boy / Kapı', f'{int(total_kasa):,}  Boy'),
        ('Pervaz (Hat 109/103)',    'Boy',  '5.0 Boy / Kapı', f'{int(total_pervaz):,}  Boy'),
    ]
    bom_gc = ['Levha', 'Seren', 'Kasa', 'Pervaz']
    for ri, (row, gc) in enumerate(zip(bom_rows, bom_gc), 5):
        for ci, val in enumerate(row, 1):
            c = ws1.cell(ri, ci, val)
            c.font = font(bold=(ci == 4), size=10)
            c.fill = fill(group_colors[gc])
            c.alignment = align('left' if ci in (1,4) else 'center')
            c.border = border_thin()
        ws1.row_dimensions[ri].height = 18

    # ── BÖLÜM 2: Hat Görev Planı ──────────────────────────────────
    ws1.merge_cells('A10:L10')
    ws1['A10'] = 'BÖLÜM 2 – HAT GÖREV DAĞILIMI & GÜNLÜK KAPASİTE (24 Saat Kesintisiz)'
    ws1['A10'].font = Font(bold=True, size=11, color='FFFFFF')
    ws1['A10'].fill = fill(C_HEADER_MED)
    ws1['A10'].alignment = align('left')
    ws1['A10'].border = border_thin()
    ws1.row_dimensions[10].height = 18

    plan_headers = [
        'Hat', 'Ürün Grubu', 'Ürün Adı', 'Boy', 'Birim',
        'Saatlik\n(boy/saat)', 'Vardiya 12s\n(boy)', 'Günlük 24s\n(boy)',
        '15.000 Kapı\nİçin Gereken',
        'Kaç Günde\nTamamlanır?',
        'Fabrika Gün\nPayı (iş günü)',
        'Notlar'
    ]
    col_widths = [10, 10, 26, 7, 7, 11, 11, 11, 14, 12, 16, 28]
    for ci, (h, w) in enumerate(zip(plan_headers, col_widths), 1):
        c = ws1.cell(11, ci, h)
        c.font = font(bold=True, size=9, color='FFFFFF')
        c.fill = fill(C_HEADER_DARK)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border_thin()
        ws1.column_dimensions[get_column_letter(ci)].width = w
    ws1.row_dimensions[11].height = 36

    # Plan verisi
    plan_data = [
        # Hat, Grup, Ürün, Boy, Birim, Saatlik, 12s, 24s, Toplam Hedef, Gün, Açıklama
        ('Hat 201', 'Levha',  'WPC Kapı Levhası (Genel Ort.)', '2.07 m', 'Plaka',
         41.7, 500, 1000, total_levha,  total_levha/1000,
         '~30 iş günü', 'Tek Hat — Tüm Levha kapasitesini karşılar'),

        ('Hat 104', 'Seren',  'Masif WPC Seren 38×45 mm', '2.10 m', 'Boy',
         28.2, 338, 676, total_seren, round(total_seren/(676+640), 1),
         '~40 iş günü', 'Ana Hat (Hat 104 + 105 kombine 40 günde kapatır)'),

        ('Hat 105', 'Seren',  'Masif WPC Seren 38×45 mm', '2.10 m', 'Boy',
         26.7, 320, 640, '–', round(total_seren/(676+640), 1),
         '~40 iş günü', 'İkinci Hat (Hat 104 ile eş zamanlı çalışır)'),

        ('Hat 107', 'Seren',  'Seren 39 mm (Yedek / 3. Hat)', '2.10 m', 'Boy',
         27.9, 335, 670, '–', round(total_seren/(676+640+670), 1),
         '~27 iş günü*', '*3 hat birlikte ~27 günde tamamlanır'),

        ('Hat 106', 'Kasa',   '140 mm Kasa', '2.15 m', 'Boy',
         14.0, 168, 336, total_kasa, round(total_kasa/(336+352+376), 1),
         '~71 iş günü', '140 mm Ana Hat'),

        ('Hat 108', 'Kasa',   '140 mm Kasa', '2.15 m', 'Boy',
         14.7, 176, 352, '–', round(total_kasa/(336+352+376), 1),
         '~71 iş günü', '140 mm İkinci Hat (Hat 106 ile eş zamanlı)'),

        ('Hat 102', 'Kasa',   '100 mm Kasa',  '2.15 m', 'Boy',
         15.7, 188, 376, '–', round(total_kasa/(336+352+376), 1),
         '~71 iş günü', '100 mm Kasa Hattı – 3. Hat olarak katkıda bulunur'),

        ('Hat 109', 'Pervaz', '50×80 Pervaz', '2.20 m', 'Boy',
         50.2, 602, 1204, total_pervaz, round(total_pervaz/(1204+756), 1),
         '~39 iş günü', 'Yüksek hızlı Pervaz Ana Hattı'),

        ('Hat 103', 'Pervaz', '50×80 Pervaz', '2.20 m', 'Boy',
         31.5, 378, 756, '–', round(total_pervaz/(1204+756), 1),
         '~39 iş günü', 'Pervaz İkinci Hat (Hat 109 ile eş zamanlı)'),
    ]

    row_i = 12
    for row in plan_data:
        hat, grup, urun, boy, birim, saatlik, v12, v24, hedef, gun, fab_gun, notlar = row
        gc = group_colors[grup]
        vals = [hat, grup, urun, boy, birim, saatlik, v12, v24, hedef, gun, fab_gun, notlar]
        for ci, val in enumerate(vals, 1):
            c = ws1.cell(row_i, ci, val)
            c.fill = fill(gc)
            c.border = border_thin()
            if ci == 1:   # Hat adı
                c.font = font(bold=True, size=10)
                c.alignment = align('center')
            elif ci == 12: # Notlar
                c.font = font(size=9, italic=True)
                c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            elif ci in (6, 7, 8, 10):
                c.font = font(bold=True, size=10)
                c.alignment = align('center')
                if isinstance(val, float):
                    c.number_format = '#,##0.0'
            elif ci == 9:
                c.font = font(bold=True, size=10)
                c.alignment = align('right')
                if isinstance(val, float):
                    c.number_format = '#,##0'
            else:
                c.font = font(size=10)
                c.alignment = align('center')
        ws1.row_dimensions[row_i].height = 22
        row_i += 1

    # ── BÖLÜM 3: Kombine Özet ─────────────────────────────────────
    row_i += 1
    ws1.merge_cells(f'A{row_i}:L{row_i}')
    ws1.cell(row_i, 1).value = 'BÖLÜM 3 – KOMBİNE TOPLAM GÜNLÜK KAPASİTE (TÜM HATLAR EŞ ZAMANLI)'
    ws1.cell(row_i, 1).font = Font(bold=True, size=11, color='FFFFFF')
    ws1.cell(row_i, 1).fill = fill(C_HEADER_MED)
    ws1.cell(row_i, 1).alignment = align('left')
    ws1.row_dimensions[row_i].height = 18
    row_i += 1

    combo_headers = ['Ürün Grubu', 'Aktif Hatlar', 'Toplam 24s Çıktı (boy/plaka)', 'Kapı Eşdeğeri / Gün', '15.000 Kapıda Kaç Gün?']
    for ci, h in enumerate(combo_headers, 1):
        c = ws1.cell(row_i, ci, h)
        c.font = font(bold=True, size=9, color='FFFFFF')
        c.fill = fill(C_HEADER_DARK)
        c.alignment = align()
        c.border = border_thin()
    ws1.row_dimensions[row_i].height = 18
    row_i += 1

    combo_data = [
        ('Levha',  'Hat 201',               1000,    1000/2.0,   total_levha/(1000/2.0)),
        ('Seren',  'Hat 104 + Hat 105',       676+640, (676+640)/3.5, total_seren/(676+640)),
        ('Seren',  'Hat 104 + 105 + 107',   676+640+670, (676+640+670)/3.5, total_seren/(676+640+670)),
        ('Kasa',   'Hat 106 + Hat 108',       336+352, (336+352)/5.0, total_kasa/(336+352)),
        ('Kasa',   'Hat 106 + 108 + Hat 102', 336+352+376, (336+352+376)/5.0, total_kasa/(336+352+376)),
        ('Pervaz', 'Hat 109 + Hat 103',      1204+756, (1204+756)/5.0, total_pervaz/(1204+756)),
    ]

    for crow in combo_data:
        grup, hatlar, daily, kapi_gun, gun_sayi = crow
        gc = group_colors[grup]
        vals = [grup, hatlar, f'{daily:,} / gün', f'{kapi_gun:,.0f} Kapı / Gün', f'{gun_sayi:.1f} iş günü']
        is_highlight = hatlar in ('Hat 104 + Hat 105', 'Hat 106 + 108 + Hat 102', 'Hat 109 + Hat 103')
        for ci, val in enumerate(vals, 1):
            c = ws1.cell(row_i, ci, val)
            c.fill = fill(C_ACCENT if is_highlight else gc)
            c.font = font(bold=is_highlight, size=10)
            c.alignment = align('center' if ci != 2 else 'left')
            c.border = border_thin()
        ws1.row_dimensions[row_i].height = 18
        row_i += 1

    # ── BÖLÜM 4: Günlük Toplam Kapasite (TAvsiye edilen hat seti) ──
    row_i += 1
    ws1.merge_cells(f'A{row_i}:L{row_i}')
    ws1.cell(row_i, 1).value = 'BÖLÜM 4 – ÖNERİLEN ROTA & TOPLAM KAPASITE ÖZETİ'
    ws1.cell(row_i, 1).font = Font(bold=True, size=11, color='FFFFFF')
    ws1.cell(row_i, 1).fill = fill(C_HEADER_MED)
    ws1.cell(row_i, 1).alignment = align('left')
    ws1.row_dimensions[row_i].height = 18
    row_i += 1

    ws1.merge_cells(f'A{row_i}:L{row_i}')
    ws1.cell(row_i, 1).value = (
        '⚙  ÖNERİLEN ROTA: Hat 201 (Levha) | Hat 104 + 105 (Seren) | '
        'Hat 106 + 108 + 102 (Kasa) | Hat 109 + 103 (Pervaz)  →  '
        'KASA DARBOĞAZdan Dolayı 15.000 Kapı ≈ 71 İş Günü  |  '
        'SEREN & PERVAZ 40 Günde Tamamlanır  →  Kasa bekleme yaratıyor.'
    )
    ws1.cell(row_i, 1).font = Font(bold=True, size=10, color='1F3864')
    ws1.cell(row_i, 1).fill = fill('FFF2CC')
    ws1.cell(row_i, 1).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws1.row_dimensions[row_i].height = 40
    row_i += 1

    # Darboğaz özeti
    darb_data = [
        ('🟢 Levha',  'Hat 201',              1000,  500,  30.0,  'En Hızlı – Levha iş günü olarak kasa bitişini bekler'),
        ('🟢 Seren',  'Hat 104 + Hat 105',    1316,  376,  40.0,  'Sereni 40 günde tamamlar, kasa bitmeden bekler'),
        ('🔴 Kasa',   'Hat 106+108+102',       1064,  213,  70.5,  'ANA DARBOĞAZ – 3 Hat eş zamanlı 71 gün'),
        ('🟢 Pervaz', 'Hat 109 + Hat 103',    1960,  392,  38.3,  'Pervazan 38 günde tamamlar, kasa bitmeden bekler'),
    ]
    darb_headers = ['Ürün Grubu', 'Hatlar', '24s Çıktı', 'Kapı Eşd./Gün', 'Süre (iş günü)', 'Yorum']
    for ci, h in enumerate(darb_headers, 1):
        c = ws1.cell(row_i, ci, h)
        c.font = font(bold=True, size=9, color='FFFFFF')
        c.fill = fill(C_HEADER_DARK)
        c.alignment = align()
        c.border = border_thin()
    ws1.row_dimensions[row_i].height = 18
    row_i += 1

    for grup, hatlar, daily, kapi_gun, gun, yorum in darb_data:
        is_red = '🔴' in grup
        gc = 'FFDCE0' if is_red else C_ACCENT
        vals = [grup, hatlar, f'{daily:,}', f'{kapi_gun:,}', f'{gun:.1f} gün', yorum]
        for ci, val in enumerate(vals, 1):
            c = ws1.cell(row_i, ci, val)
            c.fill = fill(gc)
            c.font = font(bold=is_red and ci <= 5, size=10, color='CC0000' if is_red else '000000')
            c.alignment = Alignment(horizontal='left' if ci in (1,2,6) else 'center',
                                    vertical='center', wrap_text=(ci==6))
            c.border = border_thin()
        ws1.row_dimensions[row_i].height = 22
        row_i += 1

    # ── Freeze & Filter ────────────────────────────────────────────
    ws1.freeze_panes = 'B12'

    # ═══════════════════════════════════════════════════════════════
    # SAYFA 2 – HAT GÜNLÜK DETAY ÇİZELGESİ
    # ═══════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet('Hat Günlük Detay')
    ws2.sheet_view.showGridLines = False

    ws2.merge_cells('A1:J1')
    ws2['A1'] = '15.000 KAPI – HAT GÜNLÜK KAPASİTE DETAY ÇİZELGESİ'
    ws2['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws2['A1'].fill = fill(C_HEADER_DARK)
    ws2['A1'].alignment = align()
    ws2.row_dimensions[1].height = 28

    det_headers = [
        'Hat', 'Ürün Tipi', 'Ürün Adı', 'Boy', 'Birim',
        'Saatlik\n(boy/sa)', '12s Vardiya', '24s Günlük',
        '15K İçin Gereken', 'Tamamlama\n(Gün)'
    ]
    det_widths = [10, 10, 28, 8, 7, 11, 11, 11, 18, 13]
    for ci, (h, w) in enumerate(zip(det_headers, det_widths), 1):
        c = ws2.cell(2, ci, h)
        c.font = font(bold=True, size=9, color='FFFFFF')
        c.fill = fill(C_HEADER_MED)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border_thin()
        ws2.column_dimensions[get_column_letter(ci)].width = w
    ws2.row_dimensions[2].height = 36

    # Aynı plan_data yeniden satır satır
    detail_rows = [
        ('Hat 201', 'Levha',  'WPC Kapı Levhası (Genel Ort.)', '2.07 m', 'Plaka',
         41.7, 500, 1000, int(total_levha), round(total_levha/1000, 1)),
        ('Hat 104', 'Seren',  'Masif WPC Seren 38×45 mm', '2.10 m', 'Boy',
         28.2, 338, 676, int(total_seren), round(total_seren/(676+640), 1)),
        ('Hat 105', 'Seren',  'Masif WPC Seren 38×45 mm', '2.10 m', 'Boy',
         26.7, 320, 640, '(Hat 104 ile ortak)', round(total_seren/(676+640), 1)),
        ('Hat 107', 'Seren',  'Seren 39 mm (3. Hat / Opsiyonel)', '2.10 m', 'Boy',
         27.9, 335, 670, '(3 Hat ortaklaşa)', round(total_seren/(676+640+670), 1)),
        ('Hat 106', 'Kasa',   '140 mm Kasa', '2.15 m', 'Boy',
         14.0, 168, 336, int(total_kasa), round(total_kasa/(336+352+376), 1)),
        ('Hat 108', 'Kasa',   '140 mm Kasa', '2.15 m', 'Boy',
         14.7, 176, 352, '(Hat 106 ile ortak)', round(total_kasa/(336+352+376), 1)),
        ('Hat 102', 'Kasa',   '100 mm Kasa', '2.15 m', 'Boy',
         15.7, 188, 376, '(3 Hat ortaklaşa)', round(total_kasa/(336+352+376), 1)),
        ('Hat 109', 'Pervaz', '50×80 Pervaz', '2.20 m', 'Boy',
         50.2, 602, 1204, int(total_pervaz), round(total_pervaz/(1204+756), 1)),
        ('Hat 103', 'Pervaz', '50×80 Pervaz', '2.20 m', 'Boy',
         31.5, 378, 756, '(Hat 109 ile ortak)', round(total_pervaz/(1204+756), 1)),
    ]

    zr = 3
    for row in detail_rows:
        hat, grup, urun, boy, birim, saatlik, v12, v24, hedef, gun = row
        gc = group_colors[grup]
        vals = [hat, grup, urun, boy, birim, saatlik, v12, v24, hedef, gun]
        for ci, val in enumerate(vals, 1):
            c = ws2.cell(zr, ci, val)
            c.fill = fill(gc)
            c.border = border_thin()
            c.alignment = align('left' if ci in (3,) else 'center')
            if ci == 1:
                c.font = font(bold=True)
            elif ci in (6, 7, 8):
                c.font = font(bold=True, size=10)
                if isinstance(val, (int, float)):
                    c.number_format = '#,##0'
            elif ci == 9 and isinstance(val, int):
                c.font = font(bold=True)
                c.number_format = '#,##0'
            elif ci == 10:
                c.font = font(bold=True, size=10)
                c.number_format = '#,##0.0'
            else:
                c.font = font(size=10)
        ws2.row_dimensions[zr].height = 20
        zr += 1

    ws2.auto_filter.ref = f'A2:J{zr-1}'
    ws2.freeze_panes = 'B3'

    # ─── KAYDET ───────────────────────────────────────────────────
    out_path = r'C:\Users\llboz\OneDrive\Desktop\Ergunbas_15000_Kapi_Kapasite_Plani.xlsx'
    wb.save(out_path)
    print(f'Saved: {out_path}')
    return out_path

if __name__ == '__main__':
    generate()
