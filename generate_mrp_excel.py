#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ERGUNBAS GROUP - MRP / Uretim Planlama Excel Export Scripti
Sistemdeki JSON verilerini okuyarak formatlanmis Excel raporu uretir.
Calistir: python generate_mrp_excel.py
"""

import sys
import os
import json
from datetime import datetime, date

# Windows konsol UTF-8 ayari
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("[BILGI] openpyxl yuklu degil. Yukleniyor...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

# DOSYA YOLLARI
SCRIPT_DIR   = os.path.dirname(os.path.abspath(__file__))
ORDERS_FILE  = os.path.join(SCRIPT_DIR, "orders.json")
MOLDS_FILE   = os.path.join(SCRIPT_DIR, "molds.json")
DATA_FILE    = os.path.join(SCRIPT_DIR, "data_days_2026-10.json")

# Masaustu tespiti
desktop_candidates = [
    os.path.join(os.path.expanduser("~"), "OneDrive", "Masaüstü"),
    os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop"),
    os.path.join(os.path.expanduser("~"), "Desktop"),
]
DESKTOP = next((d for d in desktop_candidates if os.path.exists(d)), SCRIPT_DIR)
OUTPUT_NAME  = f"ERGUNBAS_MRP_Uretim_Planlama_{date.today().strftime('%Y-%m-%d')}.xlsx"
OUTPUT_PATH  = os.path.join(DESKTOP, OUTPUT_NAME)

# Gun 86 = 27.10.2026
TARGET_DAY   = "86"
TARGET_DATE  = "27.10.2026"

# RENK PALETI
RENK = {
    "baslik_bg"   : "1A1A2E",   # Koyu lacivert
    "baslik_yazi" : "FFFFFF",   # Beyaz
    "alt_baslik"  : "16213E",   # Lacivert
    "alt_yazi"    : "E0E0E0",   # Acik gri
    "aktif_bg"    : "0F3460",   # Mavi
    "aktif_yazi"  : "FFFFFF",
    "idle_bg"     : "E0E0E0",   # Gri (bosta)
    "idle_yazi"   : "555555",
    "row_tek"     : "F8FAFC",   # Cok acik mavi-gri
    "row_cift"    : "FFFFFF",
    "toplam_bg"   : "E0F2FE",   # Acik mavi
    "toplam_yazi" : "0369A1",
    "uyari_bg"    : "FEF3C7",   # Sari uyari
    "uyari_yazi"  : "92400E",
    "ok_bg"       : "DCFCE7",   # Yesil OK
    "ok_yazi"     : "166534",
    "hata_bg"     : "FEE2E2",   # Kirmizi hata
    "hata_yazi"   : "991B1B",
    "border_renk" : "CBD5E1",
}

def hex2fill(hex_str, fill_type="solid"):
    return PatternFill(fill_type=fill_type, fgColor=hex_str)

def make_font(bold=False, color="000000", size=10, name="Segoe UI"):
    return Font(bold=bold, color=color, size=size, name=name)

def make_align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def make_border(all_sides="thin"):
    s = Side(style=all_sides, color=RENK["border_renk"])
    return Border(left=s, right=s, top=s, bottom=s)

def style_header_row(ws, row, cols, bg=None, fg=None, size=10, bold=True):
    bg = bg or RENK["baslik_bg"]
    fg = fg or RENK["baslik_yazi"]
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill    = hex2fill(bg)
        cell.font    = make_font(bold=bold, color=fg, size=size)
        cell.alignment = make_align("center")
        cell.border  = make_border()

def style_data_row(ws, row, cols, even=True, bg_override=None, fg_override=None, align_h="left"):
    bg = bg_override or (RENK["row_cift"] if even else RENK["row_tek"])
    fg = fg_override or "1E293B"
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        if bg_override:
            cell.fill = hex2fill(bg)
        elif not even:
            cell.fill = hex2fill(RENK["row_tek"])
        cell.font      = make_font(color=fg, size=10)
        # Sayısal değerler sağa veya ortaya
        if isinstance(cell.value, (int, float)):
            cell.alignment = make_align("right")
        else:
            cell.alignment = make_align(align_h)
        cell.border    = make_border()

def auto_width(ws, min_w=10, max_w=45):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            try:
                val = str(cell.value) if cell.value is not None else ""
                max_len = max(max_len, len(val))
            except:
                pass
        ws.column_dimensions[col_letter].width = max(min_w, min(max_len + 3, max_w))

def set_logo_row(ws, text, row=1, max_cols=12):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_cols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill      = hex2fill(RENK["baslik_bg"])
    cell.font      = make_font(bold=True, color=RENK["baslik_yazi"], size=13)
    cell.alignment = make_align("center")
    ws.row_dimensions[row].height = 28

def set_tarih_row(ws, text, row=2, max_cols=12):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_cols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill      = hex2fill(RENK["alt_baslik"])
    cell.font      = make_font(bold=False, color=RENK["alt_yazi"], size=9)
    cell.alignment = make_align("center")
    ws.row_dimensions[row].height = 18

def load_json(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        print(f"  [UYARI] Dosya bulunamadi: {path}")
        return []
    except Exception as e:
        print(f"  [HATA] {path}: {e}")
        return []

def load_day_data():
    raw = load_json(DATA_FILE)
    if not raw or not isinstance(raw, dict):
        return {}, {}
    days = raw.get("days", {})
    day  = days.get(TARGET_DAY, {})
    gunduz = day.get("gunduz", {})
    gece   = day.get("gece",   {})
    return gunduz, gece

# SHEET 1: SIPARIS PLANI
def sheet_siparis(wb, orders):
    ws = wb.create_sheet("Siparis Plani")
    ws.sheet_view.showGridLines = True

    set_logo_row(ws, "ERGUNBAS GROUP - SIPARIS & MRP PLANLAMA LISTESI", 1, 13)
    set_tarih_row(ws, f"Rapor Tarihi: {datetime.now().strftime('%d.%m.%Y %H:%M')} | Referans Uretim: {TARGET_DATE}", 2, 13)

    COLS = [
        "Siparis No", "Musteri", "Urun", "Boyut / Kesim", "Siparis Adedi",
        "Mevcut Stok", "Uretilen", "Net Ihtiyac", "Birim",
        "Atanan Hat(lar)", "Termin Tarihi", "Durum", "Not"
    ]
    ws.row_dimensions[3].height = 24
    for i, col_name in enumerate(COLS, 1):
        ws.cell(row=3, column=i, value=col_name)
    style_header_row(ws, 3, len(COLS))

    if not orders:
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(COLS))
        cell = ws.cell(row=4, column=1,
                       value="Aktif girilmis siparis kaydi bulunmamaktadir. Sisteme siparis eklendiginde bu liste otomatik dolacaktir.")
        cell.fill      = hex2fill(RENK["uyari_bg"])
        cell.font      = make_font(color=RENK["uyari_yazi"], bold=True)
        cell.alignment = make_align("center")
        ws.row_dimensions[4].height = 24
    else:
        for idx, o in enumerate(orders, 4):
            even = (idx % 2 == 0)
            sq = o.get("order_qty", 0) or 0
            st = o.get("stock_qty", 0) or 0
            pr = o.get("actual_produced", 0) or 0
            net = max(0, sq - st - pr)

            assigned = o.get("assigned_hats", [])
            assigned_str = ", ".join(str(h) for h in assigned) if isinstance(assigned, list) else str(assigned or "")

            row_data = [
                o.get("order_id", f"SIP-{idx-3}"),
                o.get("customer", "-"),
                o.get("profile_name", o.get("product", "-")),
                o.get("dim", "-"),
                sq,
                st,
                pr,
                net,
                o.get("unit", "Adet"),
                assigned_str,
                o.get("due_date", "-"),
                o.get("status", "Bekliyor"),
                o.get("note", "-"),
            ]
            for col_i, val in enumerate(row_data, 1):
                ws.cell(row=idx, column=col_i, value=val)

            status = str(o.get("status", "")).lower()
            bg_ov = None
            fg_ov = None
            if any(w in status for w in ["tamam", "bitti", "hazir"]):
                bg_ov, fg_ov = RENK["ok_bg"], RENK["ok_yazi"]
            elif any(w in status for w in ["risk", "gecik", "acil"]):
                bg_ov, fg_ov = RENK["hata_bg"], RENK["hata_yazi"]
            elif any(w in status for w in ["plan", "uretim", "calis"]):
                bg_ov, fg_ov = RENK["toplam_bg"], RENK["toplam_yazi"]

            style_data_row(ws, idx, len(COLS), even=even, bg_override=bg_ov, fg_override=fg_ov)
            ws.row_dimensions[idx].height = 20

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(COLS))}3"
    auto_width(ws)
    return ws

# SHEET 2: HAT & KALIP DURUMU
def sheet_hat_kalip(wb, molds):
    ws = wb.create_sheet("Hat & Kalip Durumu")
    ws.sheet_view.showGridLines = True

    set_logo_row(ws, f"HAT VE TAKILI KALIP DURUMU (27 Ekim 2026 Fiili Durumu)", 1, 8)
    set_tarih_row(ws, "Mevcut Takili Kaliplar ve Degisim Sureleri Tablosu", 2, 8)

    COLS = [
        "Hat No", "Takili Kalip Adi", "Kalip Kodu", "Kategori",
        "Kafa Sayisi", "Hat Durumu", "Kalip Degisim Suresi (dk)", "Aciklama / Fiili Durum"
    ]
    ws.row_dimensions[3].height = 24
    for i, col_name in enumerate(COLS, 1):
        ws.cell(row=3, column=i, value=col_name)
    style_header_row(ws, 3, len(COLS))

    status_map = {
        "active": "Aktif Uretimde",
        "idle"  : "Bosta / Beklemede",
        "maintenance": "Bakimda",
    }

    for idx, mold in enumerate(molds, 4):
        even  = (idx % 2 == 0)
        hat   = mold.get("assigned_hat", "")
        stat  = mold.get("status", "idle")
        row_data = [
            f"Hat {hat}",
            mold.get("name", ""),
            mold.get("code", ""),
            str(mold.get("category", "")).upper(),
            mold.get("heads", 1),
            status_map.get(stat, stat),
            mold.get("setup_time_min", 45),
            mold.get("notes", ""),
        ]
        for col_i, val in enumerate(row_data, 1):
            ws.cell(row=idx, column=col_i, value=val)

        if stat == "active":
            style_data_row(ws, idx, len(COLS), even=even,
                           bg_override=RENK["ok_bg"], fg_override=RENK["ok_yazi"])
        elif stat == "idle":
            style_data_row(ws, idx, len(COLS), even=even,
                           bg_override=RENK["idle_bg"], fg_override=RENK["idle_yazi"])
        else:
            style_data_row(ws, idx, len(COLS), even=even)
        ws.row_dimensions[idx].height = 20

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(COLS))}3"
    auto_width(ws)
    return ws

# SHEET 3: 27 EKIM EKSTRUDER URETIMI
def sheet_extruder(wb, gunduz, gece):
    ws = wb.create_sheet("27 Ekim Uretim Ozeti")
    ws.sheet_view.showGridLines = True

    set_logo_row(ws, f"EKSTRUDER HATLARI FIILI URETIM OZETI ({TARGET_DATE})", 1, 10)
    set_tarih_row(ws, "Gunduz ve Gece Vardiyasi Birlestirilmis Uretim Raporu", 2, 10)

    COLS = [
        "Hat No", "Uretilen Urun", "Takili Kalip",
        "Gunduz Adet", "Gunduz Kg",
        "Gece Adet", "Gece Kg",
        "TOPLAM Adet", "TOPLAM Kg",
        "Calisma Suresi (Saat)"
    ]
    ws.row_dimensions[3].height = 24
    for i, col_name in enumerate(COLS, 1):
        ws.cell(row=3, column=i, value=col_name)
    style_header_row(ws, 3, len(COLS))

    molds = load_json(MOLDS_FILE)
    kalip_by_hat = {}
    if isinstance(molds, list):
        for m in molds:
            kalip_by_hat[str(m.get("assigned_hat", ""))] = m.get("name", "")

    def extract_extruders(shift_data):
        result = {}
        for e in (shift_data.get("extruders") or []):
            hat = str(e.get("hat", ""))
            if hat not in result:
                result[hat] = {"product": e.get("product", ""), "qty": 0, "kg": 0.0, "hours": 0.0}
            result[hat]["qty"]   += e.get("qty", 0) or 0
            result[hat]["kg"]    += e.get("prod_kg", 0.0) or 0.0
            result[hat]["hours"] += e.get("hours", 0.0) or 0.0
            if e.get("product"):
                result[hat]["product"] = e["product"]
        return result

    gunduz_ext = extract_extruders(gunduz)
    gece_ext   = extract_extruders(gece)
    all_hats   = sorted(set(list(gunduz_ext.keys()) + list(gece_ext.keys())),
                        key=lambda x: int(x) if x.isdigit() else 999)

    total_adet_g = total_adet_ge = total_kg_g = total_kg_ge = 0.0
    for idx, hat in enumerate(all_hats, 4):
        even = (idx % 2 == 0)
        gd   = gunduz_ext.get(hat, {"product": "", "qty": 0, "kg": 0.0, "hours": 0.0})
        ge   = gece_ext.get(  hat, {"product": "", "qty": 0, "kg": 0.0, "hours": 0.0})
        urun = gd["product"] or ge["product"]
        kalip = kalip_by_hat.get(hat, "")
        toplam_adet = gd["qty"] + ge["qty"]
        toplam_kg   = round(gd["kg"] + ge["kg"], 1)
        toplam_saat = round(gd["hours"] + ge["hours"], 1)

        total_adet_g  += gd["qty"]
        total_adet_ge += ge["qty"]
        total_kg_g    += gd["kg"]
        total_kg_ge   += ge["kg"]

        row_data = [
            f"Hat {hat}",
            urun if urun else "[Bosta / Uretim Yok]",
            kalip,
            gd["qty"],
            round(gd["kg"], 1),
            ge["qty"],
            round(ge["kg"], 1),
            toplam_adet,
            toplam_kg,
            toplam_saat
        ]
        for col_i, val in enumerate(row_data, 1):
            ws.cell(row=idx, column=col_i, value=val)

        if not urun:
            style_data_row(ws, idx, len(COLS), even=even,
                           bg_override=RENK["idle_bg"], fg_override=RENK["idle_yazi"])
        else:
            style_data_row(ws, idx, len(COLS), even=even)
        ws.row_dimensions[idx].height = 20

    # TOPLAM satiri
    total_row = len(all_hats) + 4
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    ws.cell(row=total_row, column=1, value="GENEL TOPLAM")
    ws.cell(row=total_row, column=4, value=int(total_adet_g))
    ws.cell(row=total_row, column=5, value=round(total_kg_g, 1))
    ws.cell(row=total_row, column=6, value=int(total_adet_ge))
    ws.cell(row=total_row, column=7, value=round(total_kg_ge, 1))
    ws.cell(row=total_row, column=8, value=int(total_adet_g + total_adet_ge))
    ws.cell(row=total_row, column=9, value=round(total_kg_g + total_kg_ge, 1))
    style_header_row(ws, total_row, len(COLS),
                     bg=RENK["toplam_bg"], fg=RENK["toplam_yazi"])
    ws.row_dimensions[total_row].height = 24

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(COLS))}3"
    auto_width(ws)
    return ws

# SHEET 4: LEVHA RENK DAGILIMI
def sheet_levha(wb, gunduz, gece):
    ws = wb.create_sheet("Levha Renk Dagilimi")
    ws.sheet_view.showGridLines = True

    set_logo_row(ws, f"HAT 201 LEVHA URETIMI - RENK VE ADET DAGILIMI ({TARGET_DATE})", 1, 11)
    set_tarih_row(ws, "Mevcut Uretilen Renkler, Kesim Ebatlari ve Adetleri", 2, 11)

    COLS = [
        "Renk", "Vardiya", "En (cm)", "Boy (cm)", "Ebat Kesim",
        "Adet", "Toplam m2", "Toplam Kg",
        "Fire m2", "Fire Kg", "Calisma Saati"
    ]
    ws.row_dimensions[3].height = 24
    for i, col_name in enumerate(COLS, 1):
        ws.cell(row=3, column=i, value=col_name)
    style_header_row(ws, 3, len(COLS), bg=RENK["aktif_bg"])

    def flat_levha(shift_data, vardiya_label):
        rows = []
        levha = shift_data.get("levha", [])
        if isinstance(levha, dict):
            levha = [levha]
        for lv in (levha or []):
            rows.append({
                "renk"     : lv.get("color", ""),
                "vardiya"  : vardiya_label,
                "en"       : lv.get("width", 0),
                "boy"      : lv.get("length", 0),
                "ebat"     : f"{int(lv.get('width',0))} x {int(lv.get('length',0))} cm",
                "qty"      : lv.get("qty", 0) or 0,
                "toplam_m2": round(lv.get("total_m2", 0) or 0, 2),
                "toplam_kg": round(lv.get("total_kg", 0) or 0, 2),
                "fire_m2"  : round(lv.get("dead_fire_m2", 0) or 0, 2),
                "fire_kg"  : round(lv.get("dead_fire_kg", 0) or 0, 2),
                "hours"    : lv.get("hours", 0) or 0,
            })
        return rows

    tum_levha = flat_levha(gunduz, "Gunduz Vardiyasi") + flat_levha(gece, "Gece Vardiyasi")

    totals = {"qty": 0, "m2": 0.0, "kg": 0.0, "fire_m2": 0.0, "fire_kg": 0.0}

    for idx, lv in enumerate(tum_levha, 4):
        even = (idx % 2 == 0)
        row_data = [
            lv["renk"], lv["vardiya"], lv["en"], lv["boy"],
            lv["ebat"], lv["qty"], lv["toplam_m2"], lv["toplam_kg"],
            lv["fire_m2"], lv["fire_kg"], lv["hours"]
        ]
        for col_i, val in enumerate(row_data, 1):
            ws.cell(row=idx, column=col_i, value=val)
        style_data_row(ws, idx, len(COLS), even=even)
        ws.row_dimensions[idx].height = 20

        totals["qty"]     += lv["qty"]
        totals["m2"]      += lv["toplam_m2"]
        totals["kg"]      += lv["toplam_kg"]
        totals["fire_m2"] += lv["fire_m2"]
        totals["fire_kg"] += lv["fire_kg"]

    # Renk bazli toplanmis tablo
    sep_row = len(tum_levha) + 5
    ws.merge_cells(start_row=sep_row, start_column=1, end_row=sep_row, end_column=4)
    ws.cell(row=sep_row, column=1, value="RENK BAZLI OZET ADET TABLOSU")
    style_header_row(ws, sep_row, 4, bg=RENK["alt_baslik"])
    ws.row_dimensions[sep_row].height = 22

    ozet_header_row = sep_row + 1
    ozet_cols = ["Renk", "Toplam Adet", "Toplam m2", "Toplam Kg"]
    for ci, ch in enumerate(ozet_cols, 1):
        ws.cell(row=ozet_header_row, column=ci, value=ch)
    style_header_row(ws, ozet_header_row, 4, bg=RENK["baslik_bg"])
    ws.row_dimensions[ozet_header_row].height = 20

    renk_ozet = {}
    for lv in tum_levha:
        r = lv["renk"]
        if r not in renk_ozet:
            renk_ozet[r] = {"qty": 0, "m2": 0.0, "kg": 0.0}
        renk_ozet[r]["qty"] += lv["qty"]
        renk_ozet[r]["m2"]  += lv["toplam_m2"]
        renk_ozet[r]["kg"]  += lv["toplam_kg"]

    cur_r = ozet_header_row
    for renk, vals in renk_ozet.items():
        cur_r += 1
        row_data = [renk, vals["qty"], round(vals["m2"], 2), round(vals["kg"], 2)]
        for ci, val in enumerate(row_data, 1):
            ws.cell(row=cur_r, column=ci, value=val)
        style_data_row(ws, cur_r, 4, even=(cur_r % 2 == 0))
        ws.row_dimensions[cur_r].height = 20

    # Genel toplam satiri
    cur_r += 1
    total_data = ["GENEL TOPLAM", totals["qty"], round(totals["m2"], 2), round(totals["kg"], 2)]
    for ci, val in enumerate(total_data, 1):
        ws.cell(row=cur_r, column=ci, value=val)
    style_header_row(ws, cur_r, 4, bg=RENK["toplam_bg"], fg=RENK["toplam_yazi"])
    ws.row_dimensions[cur_r].height = 22

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(COLS))}3"
    auto_width(ws)
    return ws

# SHEET 5: KAPASITE OZETI
def sheet_kapasite(wb, gunduz, gece, molds):
    ws = wb.create_sheet("Kapasite Ozeti")
    ws.sheet_view.showGridLines = True

    now_str = datetime.now().strftime("%d.%m.%Y %H:%M")
    set_logo_row(ws, "ERGUNBAS GROUP - URETIM KAPASITE VE HAT DURUM PANOSU", 1, 6)
    set_tarih_row(ws, f"Olusturma: {now_str} | Veri Kaynagi: {TARGET_DATE} (Hat 101 - 201)", 2, 6)

    row = 3
    header_cols = ["Hat No", "Calisma Durumu", "Urun (27 Ekim)", "Takili Kalip", "Setup Suresi (dk)", "Toplam Uretim (Adet)"]
    for ci, h in enumerate(header_cols, 1):
        ws.cell(row=row, column=ci, value=h)
    style_header_row(ws, row, 6)
    ws.row_dimensions[row].height = 24

    kalip_by_hat = {}
    kalip_setup  = {}
    if isinstance(molds, list):
        for m in molds:
            h = str(m.get("assigned_hat", ""))
            kalip_by_hat[h] = m.get("name", "")
            kalip_setup[h]  = m.get("setup_time_min", 45)

    def merge_extruders(shift):
        result = {}
        for e in (shift.get("extruders") or []):
            hat = str(e.get("hat", ""))
            if hat not in result:
                result[hat] = {"product": "", "qty": 0}
            result[hat]["qty"] += e.get("qty", 0) or 0
            if e.get("product"):
                result[hat]["product"] = e["product"]
        return result

    gd_ext = merge_extruders(gunduz)
    ge_ext = merge_extruders(gece)

    hats = [str(i) for i in range(101, 110)] + ["201"]
    for idx, hat in enumerate(hats, row + 1):
        even = (idx % 2 == 0)
        gd   = gd_ext.get(hat, {"product": "", "qty": 0})
        ge   = ge_ext.get(hat, {"product": "", "qty": 0})
        urun = gd["product"] or ge["product"]
        total_qty = gd["qty"] + ge["qty"]

        if hat == "201":
            lv_total = 0
            levha_gd = gunduz.get("levha", [])
            levha_ge = gece.get("levha", [])
            if isinstance(levha_gd, dict): levha_gd = [levha_gd]
            if isinstance(levha_ge, dict): levha_ge = [levha_ge]
            for lv in (levha_gd or []) + (levha_ge or []):
                lv_total += lv.get("qty", 0) or 0
            total_qty = lv_total
            urun = "Levha (Disbudak / Koyu Ceviz / Coco)"

        durum = "Aktif Calisiyor" if urun else "Bosta / Beklemede"
        row_data = [
            f"Hat {hat}",
            durum,
            urun if urun else "-",
            kalip_by_hat.get(hat, "-"),
            kalip_setup.get(hat, "-"),
            total_qty
        ]
        for ci, val in enumerate(row_data, 1):
            ws.cell(row=idx, column=ci, value=val)

        bg_ov = RENK["ok_bg"]   if urun else RENK["idle_bg"]
        fg_ov = RENK["ok_yazi"] if urun else RENK["idle_yazi"]
        style_data_row(ws, idx, 6, even=even, bg_override=bg_ov, fg_override=fg_ov)
        ws.row_dimensions[idx].height = 20

    auto_width(ws)
    ws.freeze_panes = "A4"
    return ws

# ANA CALISTIRMA
def main():
    print("=" * 60)
    print("  ERGUNBAS GROUP - MRP Excel Export")
    print(f"  Tarih: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    print("=" * 60)

    print("\n[1/3] Veriler yukleniyor...")
    orders = load_json(ORDERS_FILE)
    if isinstance(orders, dict):
        orders = orders.get("value", orders.get("orders", []))
    molds  = load_json(MOLDS_FILE)
    if not isinstance(molds, list):
        molds = []
    gunduz, gece = load_day_data()

    print(f"  - Siparis kayitlari : {len(orders)} adet")
    print(f"  - Kalip kayitlari   : {len(molds)} adet")
    print(f"  - Gunduz extruders  : {len(gunduz.get('extruders', []))} satir")
    print(f"  - Gece extruders    : {len(gece.get('extruders', []))} satir")

    print("\n[2/3] Excel tablolari olusturuluyor...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws1 = sheet_siparis(wb, orders)
    print("  -> Sayfa 1: Siparis Plani tamamlandi")
    ws2 = sheet_hat_kalip(wb, molds)
    print("  -> Sayfa 2: Hat & Kalip Durumu tamamlandi")
    ws3 = sheet_extruder(wb, gunduz, gece)
    print("  -> Sayfa 3: 27 Ekim Uretim Ozeti tamamlandi")
    ws4 = sheet_levha(wb, gunduz, gece)
    print("  -> Sayfa 4: Levha Renk Dagilimi tamamlandi")
    ws5 = sheet_kapasite(wb, gunduz, gece, molds)
    print("  -> Sayfa 5: Kapasite Ozeti tamamlandi")

    print(f"\n[3/3] Dosya kaydediliyor...")
    wb.save(OUTPUT_PATH)
    print(f"  -> Masaustune kaydedildi: {OUTPUT_PATH}")
    print("\n" + "=" * 60)
    print("  ISLEM BASARIYLA TAMAMLANDI!")
    print("=" * 60)

if __name__ == "__main__":
    main()
